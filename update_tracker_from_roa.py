"""Fold verified Register of Actions data into the A-tier tracker."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

TRACKER = "San_Diego_ATier_Verification_Tracker.xlsx"
D = datetime.date

# Verified from the court's Register of Actions PDFs.
V = {
 "24PE001624C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,7,2), age=753,
   last=D(2026,7,16), entry="Notice of Hearing (ROA 27)",
   hearing=D(2027,3,15), hearing_re="Review Hearing re: Final Accounting",
   delinquent="YES - 2 notices (3/14/25, 5/29/26)", discharge="No",
   notes="Court issued two Notices of Failure to Perform Duties. Letters expired 4/30/2026 and were not renewed. Next review not until Mar 2027. Executor represented by an Arizona attorney."),
 "24PE001705C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,7,9), age=746,
   last=D(2026,5,27), entry="Notice of Hearing (ROA 54)",
   hearing=D(2026,7,28), hearing_re="Petition to Extend or Reissue Letters & Status of Administration",
   delinquent="YES - 2 notices (1/27/25, 4/14/26)", discharge="No",
   notes="CONTESTED. Competing petitions between Charlene Brown and Terri Gascon, plus a civil complaint alleging breach of contract, promissory estoppel, fraud and unjust enrichment. Inventory and Appraisal filed as Final $0.00. Letters expired 3/25/2026. Two law firms involved - a difficult approach."),
 "24PE000357C": dict(status="CLOSED", court_status="Case Closed - Approval Of Final Accounting", filed=D(2024,3,13), age=864,
   last=D(2026,3,5), entry="Ex Parte Petition for Final Discharge and Order (ROA 40)",
   hearing=None, hearing_re="None",
   delinquent="No", discharge="Yes",
   notes="CLOSED. Final distribution approved 1/7/2026, final discharge filed 3/5/2026. Remove from pipeline."),
 "24PE001572C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,6,28), age=757,
   last=D(2026,6,8), entry="Declaration - Withdrawal of Wells Fargo creditor claim (ROA 37)",
   hearing=D(2027,2,22), hearing_re="Review Hearing - Failure to Perform Duties re: Final Accounting",
   delinquent="YES - hearing set on it", discharge="No",
   notes="Court has set a Failure to Perform Duties review. Corrected Final Inventory and Appraisal filed 4/30/2026, so valuation work is done but distribution is not. Wells Fargo withdrew its creditor claim."),
 "24PE001744C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,7,12), age=743,
   last=D(2026,6,30), entry="Notice of Hearing (ROA 43)",
   hearing=D(2027,1,26), hearing_re="Petition for Final Distribution",
   delinquent="No", discharge="No",
   notes="WINDING DOWN. Report and Accounting plus Petition for Final Distribution filed 6/29/2026. This estate is on its way out - lowest priority of the open cases."),
 "24PE000377C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,3,14), age=863,
   last=D(2026,3,20), entry="Letters Extended/Reissued to 8/11/2026 (ROA 120)",
   hearing=D(2026,8,11), hearing_re="Accounting; also Motion to Consolidate Cases 9/2/2026",
   delinquent="YES - OSC re sanctions", discharge="No",
   notes="STRONGEST DISTRESS SIGNAL. OSC issued 3/19/2026 over the administrator distributing estate assets without court authorization, proposing $1,500 sanctions and revocation of Letters. Administrator is self-represented. Public Administrator is an interested party. 120 docket entries."),
 "24PE000954C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,5,7), age=809,
   last=D(2026,6,24), entry="Letters Extended/Reissued to 4/30/2027 (ROA 101)",
   hearing=None, hearing_re="None set",
   delinquent="No", discharge="No",
   notes="CONTESTED among family - three competing petitioners (Laurvick, Frederic Nevins, Patricia Nevins) with separate counsel. American Express is a claimant. Letters extended all the way to Apr 2027, which signals the court expects this to run long."),
 "24PE001060C": dict(status="OPEN", court_status="Pending", filed=D(2024,5,17), age=799,
   last=D(2026,2,8), entry="System Generated Notice E-mailed (ROA 28)",
   hearing=None, hearing_re="None set",
   delinquent="No", discharge="No",
   notes="Quiet since Feb 2026 - roughly five months with no filings and no hearing set. Administrator and counsel are both in St. George, Utah. Letters run to 12/11/2026."),
 "24PE000332C": dict(status="OPEN", court_status="Pending", filed=D(2024,3,12), age=865,
   last=D(2025,12,5), entry="Proof of Service (ROA 47)",
   hearing=None, hearing_re="None set",
   delinquent="No", discharge="No",
   notes="MOST DORMANT of the group - nothing filed since Dec 2025, nearly eight months, and no hearing on calendar. Letters were extended to 10/17/2027, so the administrator is planning a long runway."),
 "24PE000842C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,5,6), age=810,
   last=D(2026,7,2), entry="Review Hearing - failure to perform duties re: final accounting (ROA 26)",
   hearing=D(2026,9,17), hearing_re="Review Hearing - FAILURE TO PERFORM DUTIES, filed final accounting",
   delinquent="YES - hearing set on it", discharge="No",
   notes="Self-represented executor who is visibly struggling - two filings rejected by the clerk (1/7/2026 and 1/28/2026) and a standing Failure to Perform Duties review. No attorney of record. Best fit for a help-oriented approach."),
 "24PE001330C": dict(status="OPEN", court_status="Under Court Supervision", filed=D(2024,6,3), age=782,
   last=D(2026,4,9), entry="Accounting hearing - Amended Report and First and Final Account (ROA 28)",
   hearing=D(2026,8,25), hearing_re="Accounting; Petition for Final Distribution and fees",
   delinquent="No", discharge="No",
   notes="Moving toward final distribution, hearing 8/25/2026. Co-administrators Jane Hudson and Marion Richards with shared counsel; four additional interested parties, so consent may be slow."),
}

wb = openpyxl.load_workbook(TRACKER)
t = wb["Verify"]
hdr = [c.value for c in t[1]]
NAVY = "1F3864"

# Insert a column for the delinquency signal discovered in the dockets.
DELINQ = hdr.index("Notes") + 1          # 1-based position to insert before Notes
t.insert_cols(DELINQ)
c = t.cell(row=1, column=DELINQ, value="Court Flagged Delinquent")
c.fill = PatternFill("solid", fgColor=NAVY)
c.font = Font(bold=True, color="FFFFFF", size=10)
c.alignment = Alignment(vertical="center", wrap_text=True)
t.column_dimensions[get_column_letter(DELINQ)].width = 26
hdr = [x.value for x in t[1]]
IX = {n: i + 1 for i, n in enumerate(hdr)}

OPEN_F = PatternFill("solid", fgColor="C6EFCE")
CLOSED_F = PatternFill("solid", fgColor="F2F2F2")
FLAG_F = PatternFill("solid", fgColor="FFC7CE")
ENTRY_F = PatternFill("solid", fgColor="FFF2CC")

filled = 0
for r in range(2, t.max_row + 1):
    cn = t.cell(row=r, column=IX["Case No"]).value
    v = V.get(cn)
    if not v:
        continue
    filled += 1
    t.cell(row=r, column=IX["Status"], value=v["status"])
    t.cell(row=r, column=IX["Last Activity Date"], value=v["last"]).number_format = "yyyy-mm-dd"
    t.cell(row=r, column=IX["Last Docket Entry"], value=v["entry"])
    if v["hearing"]:
        t.cell(row=r, column=IX["Next Hearing"], value=v["hearing"]).number_format = "yyyy-mm-dd"
    t.cell(row=r, column=IX["Final Discharge?"], value=v["discharge"])
    t.cell(row=r, column=IX["Court Flagged Delinquent"], value=v["delinquent"])
    t.cell(row=r, column=IX["Notes"], value=v["notes"])

    band = OPEN_F if v["status"] == "OPEN" else CLOSED_F
    for col in ("Status", "Last Activity Date", "Last Docket Entry", "Next Hearing", "Final Discharge?"):
        t.cell(row=r, column=IX[col]).fill = band
        t.cell(row=r, column=IX[col]).font = Font(size=9)
    dc = t.cell(row=r, column=IX["Court Flagged Delinquent"])
    dc.fill = FLAG_F if v["delinquent"].startswith("YES") else band
    dc.font = Font(size=9, bold=v["delinquent"].startswith("YES"))
    nc = t.cell(row=r, column=IX["Notes"])
    nc.fill = band
    nc.font = Font(size=9)
    nc.alignment = Alignment(wrap_text=True, vertical="top")
    t.cell(row=r, column=IX["Status"]).font = Font(size=10, bold=True,
        color="006100" if v["status"] == "OPEN" else "808080")

t.column_dimensions[get_column_letter(IX["Notes"])].width = 62
t.column_dimensions[get_column_letter(IX["Last Docket Entry"])].width = 40

# Keep unverified rows visually marked as still-to-do.
for r in range(2, t.max_row + 1):
    if not t.cell(row=r, column=IX["Status"]).value:
        for col in ("Status", "Last Activity Date", "Last Docket Entry", "Next Hearing",
                    "Final Discharge?", "Court Flagged Delinquent", "Notes"):
            t.cell(row=r, column=IX[col]).fill = ENTRY_F

last = t.max_row
for dv in list(t.data_validations.dataValidation):
    t.data_validations.dataValidation.remove(dv)
d1 = DataValidation(type="list", formula1='"OPEN,CLOSED,DORMANT,NOT FOUND"', allow_blank=True)
t.add_data_validation(d1); d1.add(f"{get_column_letter(IX['Status'])}2:{get_column_letter(IX['Status'])}{last}")
d2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
t.add_data_validation(d2); d2.add(f"{get_column_letter(IX['Final Discharge?'])}2:{get_column_letter(IX['Final Discharge?'])}{last}")
t.auto_filter.ref = f"A1:{get_column_letter(t.max_column)}{last}"

# ---- Tally rebuild
q = wb["Tally"]
SC, EQ = get_column_letter(IX["Status"]), get_column_letter(IX["Equity"])
DL = get_column_letter(IX["Court Flagged Delinquent"])
for row in q.iter_rows(min_row=1, max_row=q.max_row):
    for c in row:
        c.value = None
q["A1"] = "Results"; q["A1"].font = Font(bold=True, size=14, color=NAVY)
q["A2"] = f"{filled} of {last-1} A-tier cases verified against the court docket."
q["A2"].font = Font(italic=True, size=10, color="595959")
rows = [
 ("Checked", f'=COUNTA(Verify!{SC}2:{SC}{last})'),
 ("Open", f'=COUNTIF(Verify!{SC}2:{SC}{last},"OPEN")'),
 ("Closed", f'=COUNTIF(Verify!{SC}2:{SC}{last},"CLOSED")'),
 ("Not yet checked", f'={last-1}-B4'),
 ("", ""),
 ("Open rate", f'=IF(B4=0,"",B5/B4)'),
 ("Court-flagged delinquent", f'=COUNTIF(Verify!{DL}2:{DL}{last},"YES*")'),
 ("Equity in open cases", f'=SUMIF(Verify!{SC}2:{SC}{last},"OPEN",Verify!{EQ}2:{EQ}{last})'),
 ("", ""),
 ("Projected across all 423 aged cases", ""),
 ("Est. still open", f'=IF(B4=0,"",ROUND(B9*423,0))'),
]
r = 4
for label, f in rows:
    if label:
        q.cell(row=r, column=1, value=label).font = Font(size=10, bold=label in ("Open rate", "Est. still open", "Projected across all 423 aged cases"))
        if f:
            c = q.cell(row=r, column=2, value=f)
            c.font = Font(size=10, bold=label in ("Open rate", "Est. still open"))
            if label == "Open rate": c.number_format = "0%"
            if "Equity" in label: c.number_format = '"$"#,##0'
    r += 1
q.column_dimensions["A"].width = 34
q.column_dimensions["B"].width = 18

wb.save(TRACKER)
print(f"updated {filled} rows in {TRACKER}")
PY_END = None
